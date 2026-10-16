import os
import shutil
import re
import calendar
from datetime import datetime
from copy import copy
from typing import Dict, Any, List, Tuple, Optional, Set
import openpyxl
from openpyxl.styles import Font, Alignment, Border, PatternFill

MESES_PT = [
    "janeiro", "fevereiro", "março", "marco", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"
]

MESES_NUM = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12
}

def criar_backup(caminho_planilha: str, pasta_backups: str = "backups") -> str:
    """Cria uma cópia de segurança com data e hora antes de qualquer alteração."""
    os.makedirs(pasta_backups, exist_ok=True)
    agora = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    base_nome = os.path.splitext(os.path.basename(caminho_planilha))[0]
    nome_backup = f"{base_nome} - BACKUP {agora}.xlsx"
    caminho_backup = os.path.join(pasta_backups, nome_backup)
    shutil.copy2(caminho_planilha, caminho_backup)
    return caminho_backup

def normalizar_nf(nf_val: Any) -> str:
    """Normaliza o número da NF para comparação (remove pontuação e zeros à esquerda)."""
    if nf_val is None:
        return ""
    s = str(nf_val).strip()
    s = re.sub(r"[^\d]", "", s)
    if not s:
        return ""
    try:
        return str(int(s))
    except ValueError:
        return s

def mapear_aba_produto(produto: Dict[str, Any], config_mapeamento: List[Dict[str, Any]]) -> Optional[str]:
    """
    Identifica qual aba da planilha corresponde ao produto analisando:
    1. Código do produto (05, 02, 03, etc.)
    2. Palavras-chave na descrição
    3. Palavras excluídas (ex: não classificar 'pão doce' como 'pré mistura genérica')
    """
    codigo = normalizar_nf(produto.get("codigo", ""))
    desc = produto.get("descricao", "").lower()

    for item in config_mapeamento:
        aba = item.get("aba", "").strip()
        codigos = [normalizar_nf(c) for c in item.get("codigos", [])]
        palavras = [p.lower().strip() for p in item.get("palavras_chave", [])]
        excluidas = [p.lower().strip() for p in item.get("palavras_excluidas", [])]

        if any(exc in desc for exc in excluidas):
            continue

        if codigo and codigo in codigos:
            return aba

        if any(kw in desc for kw in palavras):
            return aba

    return None

def analisar_bloco_mes_ativo(ws: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, Any]:
    """
    Localiza o bloco do mês mais recente na aba:
    - Nome do mês ativo (ex: 'Setembro')
    - Linha do cabeçalho do mês
    - Linha do saldo inicial
    - Última linha de dados preenchida
    - Próxima linha para inserção contínua
    - Conjunto de todas as NFs existentes na aba
    """
    ultimo_mes_linha = None
    ultimo_mes_nome = None

    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().lower() in MESES_PT:
            ultimo_mes_linha = r
            ultimo_mes_nome = v.strip()

    if ultimo_mes_linha is None:
        raise ValueError(f"Nenhum cabeçalho de mês encontrado na aba '{ws.title}'.")

    # Localiza a linha do saldo inicial (onde H tem valor/fórmula e A e B estão vazios)
    linha_saldo_inicial = None
    for r in range(ultimo_mes_linha, min(ultimo_mes_linha + 15, ws.max_row + 1)):
        val_h = ws.cell(r, 8).value
        val_a = ws.cell(r, 1).value
        val_b = ws.cell(r, 2).value
        eh_saldo = isinstance(val_h, (int, float)) or (isinstance(val_h, str) and val_h.startswith("="))
        if eh_saldo and val_a is None and val_b is None:
            linha_saldo_inicial = r
            break

    linha_inicio_dados = (linha_saldo_inicial + 1) if linha_saldo_inicial else (ultimo_mes_linha + 7)

    # Coleta todas as NFs da aba inteira (para prevenção global de duplicidade)
    nfs_existentes = set()
    for r in range(1, ws.max_row + 1):
        val_b = ws.cell(r, 2).value
        if val_b is not None:
            norm = normalizar_nf(val_b)
            if norm:
                nfs_existentes.add(norm)

    # Encontra a última linha que contém dados neste mês
    ultima_linha_dados = linha_saldo_inicial if linha_saldo_inicial else (linha_inicio_dados - 1)
    for r in range(linha_inicio_dados, ws.max_row + 1):
        val_b = ws.cell(r, 2).value
        val_a = ws.cell(r, 1).value
        if (val_b is not None and str(val_b).strip() != "") or (hasattr(val_a, "year") and val_a.year >= 2026):
            ultima_linha_dados = r

    # Determina a linha modelo: se já há dados no mês, usa a última linha de dados.
    # Se ainda não há dados no mês (apenas saldo inicial), usa a última linha de dados do mês anterior
    if ultima_linha_dados != linha_saldo_inicial:
        linha_modelo = ultima_linha_dados
    else:
        candidata = ultimo_mes_linha - 2
        while candidata > 1 and ws.cell(candidata, 2).value is None:
            candidata -= 1
        linha_modelo = candidata if candidata > 1 else linha_saldo_inicial

    return {
        "nome_mes": ultimo_mes_nome,
        "linha_cabecalho": ultimo_mes_linha,
        "linha_saldo_inicial": linha_saldo_inicial,
        "linha_inicio_dados": linha_inicio_dados,
        "proxima_linha": ultima_linha_dados + 1,
        "linha_modelo": linha_modelo,
        "nfs_existentes": nfs_existentes
    }

def copiar_estilo_celula(src_cell, tgt_cell):
    """Copia fonte, borda, preenchimento, alinhamento e formatação de número."""
    if src_cell.has_style:
        if src_cell.font:
            tgt_cell.font = copy(src_cell.font)
        if src_cell.border:
            tgt_cell.border = copy(src_cell.border)
        if src_cell.fill:
            tgt_cell.fill = copy(src_cell.fill)
        if src_cell.alignment:
            tgt_cell.alignment = copy(src_cell.alignment)
        if src_cell.number_format:
            tgt_cell.number_format = src_cell.number_format

def inserir_lancamento_saida(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    linha_alvo: int,
    data_obj: datetime,
    numero_nf: str,
    valor_unitario: float,
    quantidade: float,
    linha_modelo: Optional[int] = None
) -> int:
    """
    Insere o lançamento de saída na linha indicada:
    Col A: Data (datetime)
    Col B: Número NF (int se possível)
    Col C: Valor Unitário (float/int)
    Col D: Vazio
    Col E: Vazio
    Col F: Vazio
    Col G: Quantidade de Saída (float/int)
    Col H: Fórmula =H{linha_anterior}-G{linha_alvo}
    """
    try:
        nf_num = int(numero_nf)
    except ValueError:
        nf_num = numero_nf

    try:
        vu_num = float(valor_unitario)
        if vu_num.is_integer():
            vu_num = int(vu_num)
    except Exception:
        vu_num = valor_unitario

    try:
        qtd_num = float(quantidade)
        if qtd_num.is_integer():
            qtd_num = int(qtd_num)
    except Exception:
        qtd_num = quantidade

    # Garante objeto datetime para coluna de data
    if isinstance(data_obj, str):
        try:
            data_obj = datetime.strptime(data_obj[:10], "%Y-%m-%d")
        except Exception:
            pass

    # Copia formatação visual da linha modelo (fontes, bordas, cores, alinhamento)
    if linha_modelo and linha_modelo != linha_alvo:
        for c_idx in range(1, 9):
            copiar_estilo_celula(ws.cell(row=linha_modelo, column=c_idx), ws.cell(row=linha_alvo, column=c_idx))

    # Atribui os valores
    c_a = ws.cell(row=linha_alvo, column=1, value=data_obj)
    c_b = ws.cell(row=linha_alvo, column=2, value=nf_num)
    c_c = ws.cell(row=linha_alvo, column=3, value=vu_num)
    ws.cell(row=linha_alvo, column=4, value=None)
    ws.cell(row=linha_alvo, column=5, value=None)
    ws.cell(row=linha_alvo, column=6, value=None)
    c_g = ws.cell(row=linha_alvo, column=7, value=qtd_num)

    # Coluna H: se já tem fórmula mantemos, senão criamos =H{anterior}-G{atual}
    c_h = ws.cell(row=linha_alvo, column=8)
    formula_esperada = f"=H{linha_alvo - 1}-G{linha_alvo}"
    if not (isinstance(c_h.value, str) and c_h.value.startswith("=")):
        c_h.value = formula_esperada

    # Garante explicitamente os formatos de exibição corretos por coluna
    c_a.number_format = "mm-dd-yy"
    c_b.number_format = "General"
    c_c.number_format = "General"
    ws.cell(row=linha_alvo, column=4).number_format = "General"
    ws.cell(row=linha_alvo, column=5).number_format = "General"
    ws.cell(row=linha_alvo, column=6).number_format = "General"
    c_g.number_format = "General"
    c_h.number_format = '_-* #,##0_-;\\-* #,##0_-;_-* "-"??_-;_-@_-'

    return linha_alvo

def criar_cabecalho_novo_mes(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    novo_mes_nome: str,
    ano: int = 2026
) -> int:
    """
    Cria a estrutura de cabeçalho do novo mês abaixo do mês anterior:
    - Limpa linhas vazias com fórmulas antigas que ficaram no fim do mês anterior.
    - Replica exatamente a formatação, células mescladas e títulos do cabeçalho anterior.
    - Atualiza o nome do mês (ex: 'Outubro').
    - Atualiza a data do saldo anterior (ex: 'Saldo em 30/09/2026').
    - Define o saldo inicial como fórmula apontando para a última linha de saldo do mês anterior (=H{ult_linha}).
    - Cria 10 linhas preparadas com fórmula de saída.
    Retorna a linha inicial do novo cabeçalho.
    """
    info = analisar_bloco_mes_ativo(ws)
    lin_cab_orig = info['linha_cabecalho']
    lin_saldo_orig = info['linha_saldo_inicial']
    last_data_row = info['proxima_linha'] - 1

    mes_ant_nome = info['nome_mes'].lower()
    mes_ant_num = MESES_NUM.get(mes_ant_nome, 9)
    _, dia_fim = calendar.monthrange(ano, mes_ant_num)

    # Limpa linhas vazias remanescentes após last_data_row para evitar espaços em branco desnecessários
    for r in range(last_data_row + 1, ws.max_row + 1):
        for c in range(1, 14):
            ws.cell(r, c).value = None

    nova_linha_cab = last_data_row + 2
    offset = nova_linha_cab - lin_cab_orig

    # 1. Copia células, textos e estilos do cabeçalho
    for r in range(lin_cab_orig, lin_saldo_orig + 1):
        target_r = r + offset
        if ws.row_dimensions[r].height:
            ws.row_dimensions[target_r].height = ws.row_dimensions[r].height

        for c in range(1, 14):
            src = ws.cell(r, c)
            tgt = ws.cell(target_r, c)
            copiar_estilo_celula(src, tgt)

            val = src.value
            if r == lin_cab_orig and c == 1:
                tgt.value = novo_mes_nome.capitalize()
            elif isinstance(val, str) and "saldo em" in val.lower():
                tgt.value = f"Saldo em {dia_fim:02d}/{mes_ant_num:02d}/{ano}"
            elif r == lin_saldo_orig and c == 8:
                tgt.value = f"=H{last_data_row}"
            else:
                tgt.value = val

    # Garante que na linha do saldo inicial não fiquem formatos espúrios de data
    target_saldo = lin_saldo_orig + offset
    for c in range(1, 8):
        ws.cell(target_saldo, c).number_format = "General"
    ws.cell(target_saldo, 8).number_format = '_-* #,##0_-;\\-* #,##0_-;_-* "-"??_-;_-@_-'

    # 2. Replica mesclagens de células correspondentes
    merges_to_add = []
    for rng in list(ws.merged_cells.ranges):
        if lin_cab_orig <= rng.min_row <= lin_saldo_orig:
            merges_to_add.append((rng.min_row + offset, rng.max_row + offset, rng.min_col, rng.max_col))

    for min_r, max_r, min_c, max_c in merges_to_add:
        ws.merge_cells(start_row=min_r, end_row=max_r, start_column=min_c, end_column=max_c)

    # 3. Adiciona 10 linhas preparadas com fórmula de saída
    for i in range(1, 11):
        row_f = target_saldo + i
        for c in range(1, 14):
            copiar_estilo_celula(ws.cell(last_data_row, c), ws.cell(row_f, c))
            ws.cell(row_f, c).value = None
        ws.cell(row_f, 1).number_format = "mm-dd-yy"
        for c in range(2, 8):
            ws.cell(row_f, c).number_format = "General"
        ws.cell(row_f, 8).number_format = '_-* #,##0_-;\\-* #,##0_-;_-* "-"??_-;_-@_-'
        ws.cell(row_f, 8).value = f"=H{row_f-1}-G{row_f}"

    return nova_linha_cab
